#!/usr/bin/env python3
"""Pass 3: ADD new elements — missing rule numbers (S3·3) + answer keys (S3·18, S7·10, S7·23, S7·24, S8·16)."""
import copy
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

BASE = "/home/z/my-project/download/"
NAVY = RGBColor(0x0F, 0x2B, 0x4C)
GREEN = RGBColor(0x1F, 0xA3, 0x4B)

def iter_shapes(shapes):
    for sh in shapes:
        if sh.shape_type == 6:
            yield from iter_shapes(sh.shapes)
        else:
            yield sh

def styled_tb(slide, x, y, w, h, lines, align=PP_ALIGN.LEFT):
    """lines: list of list of (text, color, bold, size) run tuples per paragraph"""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, runs in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if i > 0:
            p.space_before = Pt(4)
        for text, color, bold, size in runs:
            r = p.add_run()
            r.text = text
            r.font.name = "Nunito"
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.color.rgb = color
    return tb

# ---------- 1) S3·3 clone '03' -> '01' and '02' ----------
prs = Presentation(BASE + "SESSION-3-A2-KET-TRAINER_Yola.pptx")
slide = prs.slides[2]
src = None
for sh in iter_shapes(slide.shapes):
    if sh.has_text_frame and sh.text_frame.text.strip() == "03":
        src = sh
        break
assert src is not None, "03 shape not found"
for text, x, y in (("01", 1.67, 1.86), ("02", 6.31, 1.86)):
    el = copy.deepcopy(src._element)
    slide.shapes._spTree.append(el)
    new_sh = slide.shapes[-1]
    new_sh.left = Inches(x)
    new_sh.top = Inches(y)
    for p in new_sh.text_frame.paragraphs:
        for r in p.runs:
            r.text = text
prs.save(BASE + "SESSION-3-A2-KET-TRAINER_Yola.pptx")
print("S3·3: added 01 / 02 numbers")

# ---------- 2) S3·18 Q3 answer ----------
prs = Presentation(BASE + "SESSION-3-A2-KET-TRAINER_Yola.pptx")
slide = prs.slides[17]
styled_tb(slide, 0.45, 4.78, 3.4, 0.55, [[
    ("Q3  ", NAVY, True, 11),
    ("✓ ", GREEN, True, 11),
    ("B — Jack needs Anna to make a decision.", NAVY, True, 11),
]])
prs.save(BASE + "SESSION-3-A2-KET-TRAINER_Yola.pptx")
print("S3·18: added Q3 answer")

# ---------- 3) S7·10 Cezanne answers ----------
prs = Presentation(BASE + "SESSION-7-A2-KET-TRAINER_Yola.pptx")
slide = prs.slides[9]
styled_tb(slide, 0.54, 4.12, 2.75, 1.1, [
    [("✓ Gap 1 — B: owned ", GREEN, True, 11), ("(his father had/ran the bank)", NAVY, False, 10)],
    [("✓ Gap 2 — A: career ", GREEN, True, 11), ("(a job or profession)", NAVY, False, 10)],
])
prs.save(BASE + "SESSION-7-A2-KET-TRAINER_Yola.pptx")
print("S7·10: added gap answers")

# ---------- 4) S7·23 + S7·24 exam-practice answer key (right rail) ----------
ROWS = [("19", "C earn"), ("20", "B hours"), ("21", "C worse"), ("22", "B young"), ("23", "A enough"), ("24", "B stop")]
for idx in (22, 23):
    prs = Presentation(BASE + "SESSION-7-A2-KET-TRAINER_Yola.pptx")
    slide = prs.slides[idx]
    lines = [[("ANSWERS", GREEN, True, 12)]] if False else []
    lines = [[("ANSWERS", GREEN, True, 12)]]
    for num, ans in ROWS:
        lines.append([(num + " · ", NAVY, True, 11), (ans[0] + " — ", GREEN, True, 11), (ans[2:], NAVY, False, 10.5)])
    styled_tb(slide, 8.32, 1.45, 1.45, 3.6, lines)
    prs.save(BASE + "SESSION-7-A2-KET-TRAINER_Yola.pptx")
    print(f"S7·{idx+1}: added answer key rail")

# ---------- 5) S8·16 practice answers ----------
prs = Presentation(BASE + "SESSION-8-A2-KET-TRAINER_Yola.pptx")
slide = prs.slides[15]
styled_tb(slide, 0.72, 5.06, 8.5, 0.34, [[
    ("Answers:  ", GREEN, True, 12),
    ("11 · B      12 · B      13 · A      14 · B      15 · A", NAVY, True, 12),
]])
prs.save(BASE + "SESSION-8-A2-KET-TRAINER_Yola.pptx")
print("S8·16: added answers strip")
print("DONE")
