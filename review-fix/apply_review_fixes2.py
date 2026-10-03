#!/usr/bin/env python3
"""Pass 2: fix strings that span multiple runs (character-offset replacement)."""
import json
from pptx import Presentation

BASE = "/home/z/my-project/download/"
NAVY = RGBColor = None  # not needed here

# Only the MISSED fixes from pass 1 (same tuple format)
R = []
def fx(deck, slide, old, new, label):
    R.append((deck, slide, old, new, False, label))

fx(2, 5, "You take turn to tell a story (longer than 30 seconds) that has 3 words provided",
      "You take turns to tell a story (longer than 30 seconds) using the three words given", "S2-1a")
fx(2, 5, "Each set of words must be used once only",
      "Each set of three words can be used only once", "S2-1b")
fx(2, 18, "For example: “where are you from?” Or “how old are you?”",
        "For example: “Where are you from?” / “How old are you?”", "S2-5a")
fx(2, 19, "Take turn to ask and answer questions", "Take turns to ask and answer questions", "S2-6a")
fx(2, 20, "Please choose one of the following topics and write a short paragraph (40 words max)",
        "Please choose one of the following topics. Make short notes (max 40 words), then answer aloud", "S2-7")
fx(2, 20, "Tell me about your favorite room", "Tell me about your favourite room", "S2-8")
fx(2, 22, "Watch out for distractors (wait, but,...)",
        "Listen for words like “but, actually, oh wait” — they often introduce the answer", "S2-9")
fx(4, 32, "Picture 2: She saw her friends", "Picture 2: Her friends saw her and shouted", "S4-11")
fx(5, 11, "you have to fill in the note with a single words",
        "you have to fill each gap with one word, a number, a date or a time", "S5-2")
fx(5, 20, "You and the other candidate will talk about some pictures based on the questions from the examiner.",
        "You and the other candidate will talk about five pictures. Say what you like or don’t like about them.", "S5-7a")
fx(5, 20, "After 1-2 minutes, the examiner will ask each candidate two questions related to the topic",
        "After 1-2 minutes (Phase 2), the examiner will ask each candidate two questions related to the topic", "S5-7b")
fx(5, 24, "Suggest them the best way to travel around.", "Give them advice on the best way to travel around.", "S5-9b")
fx(6, 6, "1 animal starts with letter", "one animal that starts with the letter", "S6-2a")
fx(6, 6, "1 city starts with letter", "one city that starts with the letter", "S6-2b")
fx(6, 6, "1 food starts with letter", "one food that starts with the letter", "S6-2c")
fx(6, 12, "✕ Her parents are heart doctors, not her", "✕ A — her parents are heart doctors, not her", "S6-3 A")
fx(6, 12, "✕ C. She used to, not now", "✕ C — she used to work with children, not now", "S6-3 C")
fx(7, 10, "appointment: same meaning with meeting", "appointment: similar in meaning to meeting", "S7-5")
fx(8, 5, "has the same meaning with the original one", "has the same meaning as the original one", "S8-1b")
fx(8, 12, "Important nouns, verbs, adjectives in each answer (Tim didn’t receive a message from Tom.)",
        "Important nouns, verbs, adjectives in each answer (Sarah didn’t receive a text message.)", "S8-2")
fx(8, 16, "Sarah: Oh! So who will we do, Tim?", "Sarah: Oh! So what will we do, Tim?", "S8-3a")
fx(8, 21, "You and the other candidate will talk about some pictures based on the questions from the examiner.",
        "You and the other candidate will talk about five pictures. Say what you like or don’t like about them.", "S8-7a")
fx(8, 21, "After 1-2 minutes, the examiner will ask each candidate two questions related to the topic",
        "After 1-2 minutes (Phase 2), the examiner will ask each candidate two questions related to the topic", "S8-7b")
fx(8, 30, "Be careful for distractors (but, oh wait!,...)",
        "Be careful of distractors — listen for “but, oh wait!”", "S8-11a")
fx(8, 30, "Listen to understand not to choose the answer", "Listen to understand, not just to choose the answer", "S8-11b")


def iter_shapes(shapes):
    for sh in shapes:
        if sh.shape_type == 6:
            yield from iter_shapes(sh.shapes)
        else:
            yield sh


def replace_across_runs(para, old, new):
    """Replace `old` (may span runs) with `new` inside a paragraph.
    The replacement takes the formatting of the run where `old` starts."""
    runs = para.runs
    joined = "".join(r.text for r in runs)
    idx = joined.find(old)
    if idx < 0:
        return False
    end = idx + len(old)
    pos = 0
    done_insert = False
    for r in runs:
        r_start, r_end = pos, pos + len(r.text)
        if r_end <= idx or r_start >= end:
            pos = r_end
            continue
        # overlap
        pre = r.text[: max(0, idx - r_start)]
        post = r.text[max(0, min(len(r.text), end - r_start)):]
        if not done_insert:
            r.text = pre + new + post
            done_insert = True
        else:
            r.text = pre + post
        pos = r_end
    return True


def apply():
    log = []
    for deck in sorted({r[0] for r in R}):
        fname = f"SESSION-{deck}-A2-KET-TRAINER_Yola.pptx"
        prs = Presentation(BASE + fname)
        fixes = [r for r in R if r[0] == deck]
        counts = {i: 0 for i in range(len(fixes))}
        for si, slide in enumerate(prs.slides, 1):
            for sh in iter_shapes(slide.shapes):
                frames = []
                if sh.has_text_frame:
                    frames.append(sh.text_frame)
                if getattr(sh, "has_table", False) and sh.has_table:
                    for row in sh.table.rows:
                        for cell in row.cells:
                            frames.append(cell.text_frame)
                for tf in frames:
                    for para in tf.paragraphs:
                        for k, f in enumerate(fixes):
                            _, sl, old, new, _, label = f
                            if sl is not None and sl != si:
                                continue
                            if replace_across_runs(para, old, new):
                                counts[k] += 1
        prs.save(BASE + fname)
        for k, f in enumerate(fixes):
            _, sl, old, new, _, label = f
            loc = f"S{deck}·{sl}" if sl else f"S{deck} all"
            log.append((loc, label, counts[k]))
            print(f"[{'OK ' if counts[k] else 'MISS'}] {loc:<8} {label} (x{counts[k]})")
    return log


if __name__ == "__main__":
    log = apply()
    miss = [l for l in log if l[2] == 0]
    print(f"\nPass 2 total: {len(log)} fixes, {len(miss)} still missing")
