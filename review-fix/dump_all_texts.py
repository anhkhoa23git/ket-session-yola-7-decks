#!/usr/bin/env python3
"""Dump all texts (with slide numbers) from 7 SESSION PPTX for fix planning."""
import json, sys
from pptx import Presentation

BASE = "/home/z/my-project/download/"
DECKS = {n: f"SESSION-{n}-A2-KET-TRAINER_Yola.pptx" for n in range(2, 9)}

def iter_text_frames(shapes):
    for sh in shapes:
        if sh.shape_type == 6:  # group
            yield from iter_text_frames(sh.shapes)
            continue
        if sh.has_text_frame:
            yield sh
        if getattr(sh, "has_table", False) and sh.has_table:
            for row in sh.table.rows:
                for cell in row.cells:
                    tf = cell.text_frame
                    for para in tf.paragraphs:
                        t = "".join(r.text for r in para.runs)
                        if t.strip():
                            yield ("TABLECELL", t)

out = {}
for n, fname in DECKS.items():
    prs = Presentation(BASE + fname)
    slides = []
    for i, slide in enumerate(prs.slides, 1):
        texts = []
        for item in iter_text_frames(slide.shapes):
            if isinstance(item, tuple):
                texts.append(item[1])
                continue
            for para in item.text_frame.paragraphs:
                t = "".join(r.text for r in para.runs)
                if t.strip():
                    texts.append(t)
        slides.append({"i": i, "texts": texts})
    out[n] = slides

json.dump(out, open("/home/z/my-project/work_ss/all_texts.json", "w"), ensure_ascii=False, indent=1)
print("dumped", {k: len(v) for k, v in out.items()})
